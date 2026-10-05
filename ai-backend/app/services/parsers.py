"""Trích xuất văn bản từ học liệu. Mỗi hàm trả về list[{"text", "page"?, "section"?}]."""
from pypdf import PdfReader

SUPPORTED_DOC_TYPES = ("pdf", "docx", "pptx", "txt", "md")


def extract_pdf(path: str) -> list[dict]:
    reader = PdfReader(path)
    result = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            result.append({"page": i, "text": text})
    return result


def _docx_table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        # bỏ ô trùng do merge cell
        dedup = []
        for c in cells:
            if not dedup or dedup[-1] != c:
                dedup.append(c)
        rows.append(" | ".join(dedup))
    return "\n".join(r for r in rows if r.strip(" |"))


def extract_docx(path: str) -> list[dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(path)
    sections: list[dict] = []
    heading = None
    buffer: list[str] = []

    def flush():
        if buffer:
            sections.append({"section": heading, "text": "\n".join(buffer)})
            buffer.clear()

    # duyệt theo thứ tự xuất hiện để giữ bảng đúng vị trí
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            if style.lower().startswith(("heading", "title")):
                flush()
                heading = text
            else:
                buffer.append(text)
        elif tag == "tbl":
            t = _docx_table_text(Table(child, doc))
            if t:
                buffer.append(t)
    flush()
    return sections


def extract_pptx(path: str) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(path)
    result = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        title = None
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text.strip()
                if t:
                    texts.append(t)
                    if title is None and getattr(shape, "is_placeholder", False):
                        title = t.split("\n")[0][:200]
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Ghi chú: {notes}")
        if texts:
            result.append({"page": i, "section": title, "text": "\n".join(texts)})
    return result


def extract_text_file(path: str) -> list[dict]:
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            content = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return extract_plain_text(content)


def extract_plain_text(content: str) -> list[dict]:
    """Tách theo tiêu đề markdown (#) làm section."""
    sections: list[dict] = []
    heading = None
    buffer: list[str] = []
    for line in content.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            if buffer and "".join(buffer).strip():
                sections.append({"section": heading, "text": "\n".join(buffer)})
            buffer = []
            heading = stripped.lstrip("#").strip() or heading
        else:
            buffer.append(line)
    if buffer and "".join(buffer).strip():
        sections.append({"section": heading, "text": "\n".join(buffer)})
    return sections


def extract(path: str, source_type: str) -> list[dict]:
    if source_type == "pdf":
        return extract_pdf(path)
    if source_type == "docx":
        return extract_docx(path)
    if source_type == "pptx":
        return extract_pptx(path)
    if source_type in ("txt", "md"):
        return extract_text_file(path)
    raise ValueError(f"Định dạng không hỗ trợ: {source_type}")
