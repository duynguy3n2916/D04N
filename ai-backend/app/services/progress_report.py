"""Báo cáo tiến độ học tập của học sinh (DOCX): XP, chuỗi ngày học, tiến độ từng khóa và kết quả bài kiểm tra."""
import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.core.security import Principal
from app.models.accounts import User
from app.models.learning import ItemProgress, LessonItem
from app.services import learning


def collect(db: Session, user: Principal) -> dict:
    student = db.get(User, user.user_id)
    courses = [learning.course_tree(db, user, c) for c in learning.visible_courses(db, user)]
    quizzes = []
    rows = (db.query(ItemProgress, LessonItem).join(LessonItem, ItemProgress.item_id == LessonItem.id)
            .filter(ItemProgress.student_id == user.user_id, LessonItem.type == "quiz")
            .order_by(LessonItem.title).all())
    for prog, item in rows:
        meta = prog.meta or {}
        quizzes.append({
            "title": item.title, "lesson": item.lesson.code, "course": item.lesson.chapter.course.title,
            "best": round(prog.score * 100) if prog.score is not None else None,
            "passed": prog.status == "completed", "finished_at": meta.get("finished_at"),
        })
    return {"name": (student.display_name if student else None) or user.user_id, "user_id": user.user_id,
            "stats": learning.stats(db, user), "courses": courses, "quizzes": quizzes, "today": date.today()}


def to_docx(data: dict) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_run("BÁO CÁO TIẾN ĐỘ HỌC TẬP").bold = True
    sub = doc.add_paragraph(f"Học viên: {data['name']} (@{data['user_id']}) · Ngày lập: {data['today'].strftime('%d/%m/%Y')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    s = data["stats"]
    doc.add_heading("1. Tổng quan", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Chỉ số", "Giá trị"
    for label, value in [("Tổng XP", s["total_xp"]), ("XP trong tuần", s["week_xp"]),
                         ("Chuỗi ngày học liên tiếp", f"{s['streak_days']} ngày"),
                         ("Mục tiêu XP mỗi ngày", s["daily_goal_xp"])]:
        r = t.add_row().cells
        r[0].text, r[1].text = label, str(value)

    doc.add_heading("2. Tiến độ các khóa học", level=2)
    if not data["courses"]:
        doc.add_paragraph("Chưa có khóa học nào.")
    for c in data["courses"]:
        total, done = c["items_total"], c["items_done"]
        pct = round(done / total * 100) if total else 0
        doc.add_paragraph(f"{c['title']}: hoàn thành {done}/{total} mục ({pct}%)", style="List Bullet")

    doc.add_heading("3. Bài kiểm tra", level=2)
    if not data["quizzes"]:
        doc.add_paragraph("Chưa làm bài kiểm tra nào.")
    else:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for cell, text in zip(t.rows[0].cells, ["Bài kiểm tra", "Khóa / bài", "Điểm cao nhất", "Kết quả"]):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for q in data["quizzes"]:
            r = t.add_row().cells
            r[0].text, r[1].text = q["title"], f"{q['course']} · {q['lesson']}"
            r[2].text = f"{q['best']}%" if q["best"] is not None else "—"
            r[3].text = "Đạt" if q["passed"] else "Chưa đạt"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
