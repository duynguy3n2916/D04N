"""Xuất đề kiểm tra từ ngân hàng câu hỏi: DOCX (python-docx) và HTML để in / lưu PDF từ trình duyệt.

- Nhiều mã đề: mỗi mã xáo thứ tự câu và thứ tự lựa chọn (tất định theo mã, xuất lại vẫn giống), đáp án tính lại theo mã.
- Phần đáp án + giải thích để ở trang riêng cuối file (tùy chọn).
"""
import html
import io
import random
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from app.core.utils import strip_option_label

LETTERS = "ABCDEFGH"
TYPE_LABEL = {"multiple_choice": "Trắc nghiệm", "true_false": "Đúng / Sai", "short_answer": "Trả lời ngắn",
              "essay": "Tự luận", "open": "Tự luận"}


@dataclass
class ExamQuestion:
    question: str
    type: str
    options: list[str] | None  # nội dung lựa chọn đã bỏ nhãn "A. "
    correct_index: int | None
    answer: str  # đáp án dạng chữ (câu tự luận: đáp án mẫu)
    explanation: str


def _from_payload(p: dict) -> ExamQuestion:
    options = [strip_option_label(o) for o in (p.get("options") or [])] or None
    idx = p.get("correct_index")
    if options and (idx is None or not 0 <= idx < len(options)):
        idx = None
    answer = options[idx] if options and idx is not None else (p.get("correct_answer") or "")
    return ExamQuestion(question=(p.get("question") or "").strip(), type=p.get("type") or "multiple_choice",
                        options=options, correct_index=idx, answer=answer, explanation=(p.get("explanation") or "").strip())


def build_variant(payloads: list[dict], variant: int, shuffle: bool) -> list[ExamQuestion]:
    qs = [_from_payload(p) for p in payloads]
    if not shuffle:
        return qs
    rnd = random.Random(f"de-{variant}")
    rnd.shuffle(qs)
    for q in qs:
        # câu Đúng/Sai giữ thứ tự tự nhiên; trắc nghiệm xáo lựa chọn và tính lại vị trí đáp án đúng
        if q.type == "multiple_choice" and q.options and q.correct_index is not None:
            order = list(range(len(q.options)))
            rnd.shuffle(order)
            q.options = [q.options[i] for i in order]
            q.correct_index = order.index(q.correct_index)
    return qs


def _key(q: ExamQuestion) -> str:
    if q.options and q.correct_index is not None:
        return LETTERS[q.correct_index] if q.type != "true_false" else q.options[q.correct_index]
    return q.answer or "—"


def _variant_code(v: int) -> str:
    return f"{100 + v}"


def to_docx(title: str, subtitle: str, variants: list[list[ExamQuestion]], include_answers: bool,
            duration_minutes: int | None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    for vi, qs in enumerate(variants, 1):
        if vi > 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(title.upper())
        r.bold, r.font.size = True, Pt(15)
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info = [subtitle] if subtitle else []
        if duration_minutes:
            info.append(f"Thời gian làm bài: {duration_minutes} phút")
        if len(variants) > 1:
            info.append(f"Mã đề: {_variant_code(vi)}")
        meta.add_run(" · ".join(info)).italic = True
        doc.add_paragraph("Họ và tên: ........................................................   Lớp: ....................")

        for n, q in enumerate(qs, 1):
            p = doc.add_paragraph()
            p.add_run(f"Câu {n}. ").bold = True
            p.add_run(q.question)
            if q.options:
                for i, opt in enumerate(q.options):
                    doc.add_paragraph(f"{LETTERS[i]}. {opt}").paragraph_format.left_indent = Pt(18)
            else:
                lines = 6 if q.type in ("essay", "open") else 2
                for _ in range(lines):
                    doc.add_paragraph("." * 110)

    if include_answers:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run("ĐÁP ÁN VÀ HƯỚNG DẪN")
        r.bold, r.font.size = True, Pt(15)
        for vi, qs in enumerate(variants, 1):
            if len(variants) > 1:
                doc.add_paragraph().add_run(f"Mã đề {_variant_code(vi)}").bold = True
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ("Câu", "Đáp án", "Giải thích")):
                cell.paragraphs[0].add_run(text).bold = True
            for n, q in enumerate(qs, 1):
                row = table.add_row().cells
                row[0].text, row[1].text, row[2].text = str(n), _key(q), q.explanation
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_html(title: str, subtitle: str, variants: list[list[ExamQuestion]], include_answers: bool,
            duration_minutes: int | None) -> str:
    e = html.escape
    parts = [f"""<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>{e(title)}</title>
<style>
body{{font-family:"Times New Roman",serif;font-size:13pt;max-width:780px;margin:24px auto;padding:0 16px;color:#111;background:#fff}}
h1{{text-align:center;font-size:16pt;margin:0 0 4px}} .meta{{text-align:center;font-style:italic;margin-bottom:12px}}
.q{{margin:10px 0;break-inside:avoid}} .opt{{margin-left:22px}} .line{{border-bottom:1px dotted #555;height:22px}}
table{{border-collapse:collapse;width:100%}} td,th{{border:1px solid #333;padding:4px 6px;text-align:left;vertical-align:top}}
.page{{break-after:page}} @media print{{body{{margin:0}} .noprint{{display:none}}}}
</style></head><body>
<p class="noprint" style="text-align:right"><button onclick="window.print()">In / Lưu PDF</button></p>"""]
    for vi, qs in enumerate(variants, 1):
        info = [subtitle] if subtitle else []
        if duration_minutes:
            info.append(f"Thời gian làm bài: {duration_minutes} phút")
        if len(variants) > 1:
            info.append(f"Mã đề: {_variant_code(vi)}")
        last = vi == len(variants) and not include_answers
        parts.append(f'<section class="{"" if last else "page"}"><h1>{e(title.upper())}</h1>'
                     f'<div class="meta">{e(" · ".join(info))}</div>'
                     "<p>Họ và tên: ....................................................... Lớp: ..................</p>")
        for n, q in enumerate(qs, 1):
            parts.append(f'<div class="q"><b>Câu {n}.</b> {e(q.question)}')
            if q.options:
                parts += [f'<div class="opt">{LETTERS[i]}. {e(o)}</div>' for i, o in enumerate(q.options)]
            else:
                parts += ['<div class="line"></div>'] * (6 if q.type in ("essay", "open") else 2)
            parts.append("</div>")
        parts.append("</section>")
    if include_answers:
        parts.append("<section><h1>ĐÁP ÁN VÀ HƯỚNG DẪN</h1>")
        for vi, qs in enumerate(variants, 1):
            if len(variants) > 1:
                parts.append(f"<h3>Mã đề {_variant_code(vi)}</h3>")
            parts.append("<table><tr><th>Câu</th><th>Đáp án</th><th>Giải thích</th></tr>")
            parts += [f"<tr><td>{n}</td><td>{e(_key(q))}</td><td>{e(q.explanation)}</td></tr>" for n, q in enumerate(qs, 1)]
            parts.append("</table>")
        parts.append("</section>")
    parts.append("</body></html>")
    return "".join(parts)


def export(payloads: list[dict], *, title: str, subtitle: str = "", fmt: str = "docx", include_answers: bool = True,
           variants: int = 1, shuffle: bool = False, duration_minutes: int | None = None) -> tuple[bytes, str, str]:
    """Trả (nội dung, media type, phần mở rộng file)."""
    shuffle = shuffle or variants > 1
    sets = [build_variant(payloads, v, shuffle) for v in range(1, variants + 1)]
    subtitle = subtitle or f"Ngày {date.today().strftime('%d/%m/%Y')} · {len(payloads)} câu"
    if fmt == "html":
        return to_html(title, subtitle, sets, include_answers, duration_minutes).encode("utf-8"), \
            "text/html; charset=utf-8", "html"
    return to_docx(title, subtitle, sets, include_answers, duration_minutes), \
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
